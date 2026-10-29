import argparse
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from audio_pipeline.scripts.db import (
    get_connection,
    get_import_job,
)


REPORT_DIR = Path("/opt/musicapp/reports")


def parse_arguments():
    parser = argparse.ArgumentParser(
        description=(
            "Generate an Excel report for music "
            "import tracking records."
        )
    )

    parser.add_argument(
        "--job-id",
        type=int,
        default=None,
        help=(
            "Generate the report only for this "
            "import job."
        ),
    )

    return parser.parse_args()


def get_report_items(job_id=None):
    connection = get_connection()

    try:
        with connection.cursor() as cursor:
            if job_id is None:
                cursor.execute(
                    """
                    SELECT
                        i.id,
                        i.job_id,
                        i.drive_file_id,
                        i.original_filename,
                        i.title,
                        i.artist,
                        i.album,
                        i.status,
                        i.playback_validated,
                        i.song_id,
                        i.sha256,
                        i.local_file_path,
                        i.failure_reason,
                        i.retry_count,
                        i.created_at,
                        i.updated_at
                    FROM import_items i
                    ORDER BY i.id;
                    """
                )
            else:
                cursor.execute(
                    """
                    SELECT
                        i.id,
                        i.job_id,
                        i.drive_file_id,
                        i.original_filename,
                        i.title,
                        i.artist,
                        i.album,
                        i.status,
                        i.playback_validated,
                        i.song_id,
                        i.sha256,
                        i.local_file_path,
                        i.failure_reason,
                        i.retry_count,
                        i.created_at,
                        i.updated_at
                    FROM import_items i
                    WHERE i.job_id = %s
                    ORDER BY i.id;
                    """,
                    (job_id,),
                )

            return cursor.fetchall()

    finally:
        connection.close()


def calculate_counts(rows):
    validated = sum(
        1
        for row in rows
        if row[7] == "validated"
    )

    warnings = sum(
        1
        for row in rows
        if row[7] == "warning"
    )

    duplicates = sum(
        1
        for row in rows
        if row[7] == "duplicate"
    )

    failed = sum(
        1
        for row in rows
        if row[7] == "failed"
    )

    playback_passed = sum(
        1
        for row in rows
        if row[8] is True
    )

    successful_imports = sum(
        1
        for row in rows
        if row[7] in (
            "validated",
            "warning",
        )
        and row[8] is True
        and row[9] is not None
    )

    return {
        "total": len(rows),
        "successful": successful_imports,
        "validated": validated,
        "warnings": warnings,
        "duplicates": duplicates,
        "failed": failed,
        "playback_passed": playback_passed,
    }


def create_summary_sheet(
    workbook,
    rows,
    generated_at,
    job=None,
):
    sheet = workbook.active
    sheet.title = "Summary"

    counts = calculate_counts(rows)

    summary = [
        (
            "Madineni Music Import Report",
            "",
        ),
        (
            "Generated At (UTC)",
            generated_at,
        ),
    ]

    if job is not None:
        summary.extend(
            [
                (
                    "Import Job ID",
                    job["id"],
                ),
                (
                    "Import Run ID",
                    job["run_id"],
                ),
                (
                    "Job Status",
                    job["status"],
                ),
            ]
        )

    summary.extend(
        [
            ("", ""),
            (
                "Total Tracked Files",
                counts["total"],
            ),
            (
                "Successful Imports",
                counts["successful"],
            ),
            (
                "Validated",
                counts["validated"],
            ),
            (
                "Warnings",
                counts["warnings"],
            ),
            (
                "Duplicates",
                counts["duplicates"],
            ),
            (
                "Failed",
                counts["failed"],
            ),
            (
                "Playback Validated",
                counts["playback_passed"],
            ),
        ]
    )

    for row_number, values in enumerate(
        summary,
        start=1,
    ):
        sheet.cell(
            row=row_number,
            column=1,
            value=values[0],
        )

        sheet.cell(
            row=row_number,
            column=2,
            value=values[1],
        )

    sheet["A1"].font = Font(
        bold=True,
        size=16,
    )

    for row_number in range(
        1,
        len(summary) + 1,
    ):
        label = sheet.cell(
            row=row_number,
            column=1,
        ).value

        if label:
            sheet.cell(
                row=row_number,
                column=1,
            ).font = Font(
                bold=True
            )

    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 48


def create_detail_sheet(
    workbook,
    rows,
):
    sheet = workbook.create_sheet(
        "Import Report"
    )

    headers = [
        "Import Item ID",
        "Job ID",
        "Drive File ID",
        "Original Filename",
        "Title",
        "Artist",
        "Album",
        "Status",
        "Playback Validated",
        "Song ID",
        "SHA256",
        "Local File Path",
        "Failure Reason",
        "Retry Count",
        "Created At",
        "Updated At",
    ]

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    for column_number, header in enumerate(
        headers,
        start=1,
    ):
        cell = sheet.cell(
            row=1,
            column=column_number,
            value=header,
        )

        cell.font = Font(
            bold=True
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for row_number, row in enumerate(
        rows,
        start=2,
    ):
        for (
            column_number,
            value,
        ) in enumerate(
            row,
            start=1,
        ):
            if isinstance(
                value,
                datetime,
            ):
                value = value.isoformat()

            sheet.cell(
                row=row_number,
                column=column_number,
                value=value,
            )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    widths = {
        1: 15,
        2: 12,
        3: 32,
        4: 42,
        5: 32,
        6: 38,
        7: 28,
        8: 14,
        9: 20,
        10: 12,
        11: 68,
        12: 70,
        13: 50,
        14: 12,
        15: 28,
        16: 28,
    }

    for (
        column_number,
        width,
    ) in widths.items():
        sheet.column_dimensions[
            get_column_letter(
                column_number
            )
        ].width = width


def generate_report(job_id=None):
    job = None

    if job_id is not None:
        job = get_import_job(job_id)

        if job is None:
            raise RuntimeError(
                f"Import job {job_id} "
                "does not exist."
            )

    rows = get_report_items(
        job_id=job_id
    )

    # A valid import job may contain zero items when Drive INCOMING is empty.
    # Keep generating the standard workbook with zero counts and headers.
    # For unscoped/manual reports, retain the original no-data safeguard.
    if not rows and job_id is None:
        raise RuntimeError(
            "No import_items found. "
            "Report was not created."
        )

    REPORT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    now = datetime.now(
        timezone.utc
    )

    if job_id is not None:
        report_path = REPORT_DIR / (
            "music_import_report_"
            f"job_{job_id}.xlsx"
        )
    else:
        timestamp = now.strftime(
            "%Y%m%d_%H%M%S"
        )

        report_path = REPORT_DIR / (
            "music_import_report_"
            f"{timestamp}.xlsx"
        )

    workbook = Workbook()

    create_summary_sheet(
        workbook,
        rows,
        now.isoformat(),
        job=job,
    )

    create_detail_sheet(
        workbook,
        rows,
    )

    workbook.save(
        report_path
    )

    if not report_path.is_file():
        raise RuntimeError(
            "Report file was not created."
        )

    if (
        report_path.stat().st_size
        == 0
    ):
        raise RuntimeError(
            "Report file is empty."
        )

    counts = calculate_counts(
        rows
    )

    print("=" * 70)
    print(
        "EXCEL IMPORT REPORT"
    )
    print("=" * 70)

    if job is not None:
        print(
            f"Import job ID:       "
            f"{job['id']}"
        )
        print(
            f"Import run:          "
            f"{job['run_id']}"
        )

    print(
        f"Report:              "
        f"{report_path}"
    )
    print(
        f"Tracked files:       "
        f"{counts['total']}"
    )
    print(
        f"Successful imports:  "
        f"{counts['successful']}"
    )
    print(
        f"Warnings:            "
        f"{counts['warnings']}"
    )
    print(
        f"Duplicates:          "
        f"{counts['duplicates']}"
    )
    print(
        f"Failed:              "
        f"{counts['failed']}"
    )
    print(
        f"Playback validated:  "
        f"{counts['playback_passed']}"
    )
    print(
        f"Size:                "
        f"{report_path.stat().st_size} "
        "bytes"
    )
    print()
    print(
        "Excel report created successfully."
    )

    return report_path


def main():
    args = parse_arguments()

    generate_report(
        job_id=args.job_id
    )


if __name__ == "__main__":
    main()