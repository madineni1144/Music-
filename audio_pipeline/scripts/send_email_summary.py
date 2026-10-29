import argparse
import os
import smtplib
from email.message import EmailMessage
from pathlib import Path

from dotenv import dotenv_values
from openpyxl import load_workbook

from audio_pipeline.scripts.db import get_import_job


EMAIL_ENV = Path(
    os.getenv(
        "MUSICAPP_EMAIL_ENV",
        "/etc/musicapp-email.env",
    )
)

REPORT_DIR = Path("/opt/musicapp/reports")
SMTP_HOST = "smtp.gmail.com"
SMTP_PORT = 465


def parse_arguments():
    parser = argparse.ArgumentParser(
        description="Send music import email summary."
    )
    parser.add_argument(
        "--job-id",
        type=int,
        default=None,
        help="Send the report for this import job.",
    )
    return parser.parse_args()


def load_email_config():
    if not EMAIL_ENV.is_file():
        raise FileNotFoundError(
            f"Email configuration not found: {EMAIL_ENV}"
        )

    config = dotenv_values(EMAIL_ENV)

    address = config.get("GMAIL_ADDRESS")
    password = config.get("GMAIL_APP_PASSWORD")
    recipient = config.get("EMAIL_TO")

    if not all((address, password, recipient)):
        raise RuntimeError(
            "Missing GMAIL_ADDRESS, "
            "GMAIL_APP_PASSWORD, or EMAIL_TO."
        )

    return address, password.replace(" ", ""), recipient


def get_report(job_id=None):
    if job_id is not None:
        path = REPORT_DIR / (
            f"music_import_report_job_{job_id}.xlsx"
        )
        if not path.is_file():
            raise FileNotFoundError(path)
        return path

    reports = list(
        REPORT_DIR.glob("music_import_report_*.xlsx")
    )

    if not reports:
        raise FileNotFoundError("No Excel reports found.")

    return max(
        reports,
        key=lambda path: path.stat().st_mtime,
    )


def read_summary(report_path):
    workbook = load_workbook(
        report_path,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook["Summary"]
        return {
            str(key): value
            for key, value in sheet.iter_rows(
                min_row=1,
                max_col=2,
                values_only=True,
            )
            if key
        }
    finally:
        workbook.close()


def send_email(job_id=None):
    job = None

    if job_id is not None:
        job = get_import_job(job_id)

        if job is None:
            raise RuntimeError(
                f"Import job {job_id} does not exist."
            )

    report_path = get_report(job_id)

    marker = REPORT_DIR / (
        f".email_sent_job_{job_id}.marker"
    ) if job_id is not None else None

    if marker is not None and marker.exists():
        print(
            f"SKIPPED: Email already recorded "
            f"as sent for Job {job_id}."
        )
        return

    summary = read_summary(report_path)

    if job_id is not None:
        report_job_id = summary.get("Import Job ID")

        if str(report_job_id) != str(job_id):
            raise RuntimeError(
                "Excel report Job ID does not match "
                "requested Job ID."
            )

    address, password, recipient = load_email_config()

    total = summary.get("Total Tracked Files", 0)
    successful = summary.get("Successful Imports", 0)
    validated = summary.get("Validated", 0)
    warnings = summary.get("Warnings", 0)
    duplicates = summary.get("Duplicates", 0)
    failed = summary.get("Failed", 0)
    playback = summary.get("Playback Validated", 0)
    generated_at = summary.get(
        "Generated At (UTC)", "Unknown"
    )

    message = EmailMessage()
    message["From"] = address
    message["To"] = recipient

    job_label = (
        f"Job {job_id} - "
        if job_id is not None else ""
    )

    message["Subject"] = (
        f"Madineni Music Import - {job_label}"
        f"{successful} Successful / {failed} Failed"
    )

    message.set_content(
        f"""Madineni Music Import Pipeline Summary

Import Job:          {job_id if job_id is not None else 'All'}
Generated At (UTC):  {generated_at}

Total Tracked Files: {total}
Successful Imports:  {successful}
Validated:           {validated}
Warnings:            {warnings}
Duplicates:          {duplicates}
Failed:              {failed}
Playback Validated:  {playback}

Excel Report:
{report_path.name}

The Excel report is attached.

Madineni Music Pipeline
"""
    )

    message.add_attachment(
        report_path.read_bytes(),
        maintype="application",
        subtype=(
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        filename=report_path.name,
    )

    with smtplib.SMTP_SSL(
        SMTP_HOST,
        SMTP_PORT,
        timeout=30,
    ) as smtp:
        smtp.login(address, password)
        smtp.send_message(message)

    if marker is not None:
        with marker.open("x") as file:
            file.write(
                f"job_id={job_id}\n"
                f"report={report_path.name}\n"
            )

    print("=" * 70)
    print("EMAIL SUMMARY")
    print("=" * 70)

    if job is not None:
        print(f"Import job ID:       {job['id']}")
        print(f"Import run:          {job['run_id']}")

    print(f"From:                {address}")
    print(f"To:                  {recipient}")
    print(f"Report:              {report_path.name}")
    print(f"Tracked files:       {total}")
    print(f"Successful imports:  {successful}")
    print(f"Warnings:            {warnings}")
    print(f"Duplicates:          {duplicates}")
    print(f"Failed:              {failed}")
    print(f"Playback validated:  {playback}")
    print()
    print("Email summary sent successfully.")


def main():
    args = parse_arguments()
    send_email(job_id=args.job_id)


if __name__ == "__main__":
    main()
